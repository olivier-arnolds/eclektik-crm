#!/usr/bin/env python3
"""
Bouwt een Excel-totaaloverzicht van de Glint-prospectmarkt.

WAAROM DIT SCRIPT
  Het overzicht komt uit drie bronnen die los van elkaar leven: de accounts die in
  het CRM als 'Expected Glint Customers' staan, de marktanalyse van de collega
  (xlsx, staat niet in het CRM) en de Surfe-validatieronde. Los van elkaar zeggen
  ze weinig; naast elkaar laten ze zien wie we vandaag kunnen benaderen en waar
  het aan schort. Draaien kan opnieuw zodra een van de drie bijwerkt.

GEBRUIK
  python3 scripts/export-glint-prospects.py [--out /pad/naar/bestand.xlsx]

  Leest SUPABASE_URL + SUPABASE_SERVICE_KEY uit .env.local in de repo-root.
  Het script schrijft NIETS naar de database; het leest alleen.
"""
import argparse, csv, json, os, re, sys, unicodedata
import urllib.parse, urllib.request
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ANALYSE_XLSX = os.path.expanduser("~/Downloads/Eclectik_Viva_Glint_Prospect.xlsx")
SURFE_IN = os.path.expanduser("~/Downloads/surfe-input-glint-prioriteit-a.csv")
SURFE_UIT = os.path.expanduser("~/Downloads/surfe-input-glint-prioriteit-a (1).csv")

GLINT_TYPE = "Expected Glint Customers"


def env():
    pad = os.path.join(ROOT, ".env.local")
    waarden = {}
    with open(pad) as f:
        for regel in f:
            if "=" in regel and not regel.strip().startswith("#"):
                k, _, v = regel.partition("=")
                waarden[k.strip()] = v.strip().strip('"')
    url, key = waarden.get("SUPABASE_URL"), waarden.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit("SUPABASE_URL of SUPABASE_SERVICE_KEY ontbreekt in .env.local")
    return url.rstrip("/"), key


def haal(url, key, tabel, kolommen, filter_=""):
    """Alles ophalen in pagina's. PostgREST geeft standaard maximaal 1000 rijen
    terug, dus zonder paginering mis je stilzwijgend de rest."""
    uit, offset = [], 0
    while True:
        pad = (f"{url}/rest/v1/{tabel}?select={urllib.parse.quote(kolommen)}"
               f"&limit=1000&offset={offset}{filter_}")
        req = urllib.request.Request(pad, headers={
            "apikey": key, "Authorization": f"Bearer {key}", "Accept": "application/json"})
        with urllib.request.urlopen(req, timeout=120) as r:
            blok = json.loads(r.read().decode())
        uit.extend(blok)
        if len(blok) < 1000:
            return uit
        offset += 1000


def norm(s):
    s = unicodedata.normalize("NFKD", (s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", s.lower())


def benaderbaar(c):
    """Kan deze persoon vandaag een mail krijgen? Elke reden apart, want 'nee'
    zonder reden is in een overzicht net zo onbruikbaar als geen overzicht."""
    if not (c.get("email") or "").strip():
        return "nee", "geen e-mailadres"
    if c.get("do_not_email"):
        return "nee", "do_not_email"
    if c.get("former"):
        return "nee", "former employee"
    if (c.get("stage") or "").lower() == "inactive":
        return "nee", "inactief"
    return "ja", ""


def kop(ws, kolommen):
    vul = PatternFill("solid", fgColor="1F3864")
    for i, naam in enumerate(kolommen, 1):
        cel = ws.cell(row=1, column=i, value=naam)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = vul
        cel.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(kolommen))}1"


def breedtes(ws, kolommen, rijen):
    for i, naam in enumerate(kolommen, 1):
        langste = max([len(str(naam))] + [len(str(r[i - 1] or "")) for r in rijen[:400]])
        ws.column_dimensions[get_column_letter(i)].width = min(max(langste + 2, 10), 46)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--out", default=os.path.expanduser(
        f"~/Downloads/glint-prospects-totaaloverzicht-{date.today():%Y-%m-%d}.xlsx"))
    args = p.parse_args()
    url, key = env()

    bedrijven = haal(url, key, "companies",
                     "id,name,account_no,industry,sub_industry,country,city,employee_count,"
                     "website,linkedin_url,stage,type,owner,notes",
                     f"&type=eq.{urllib.parse.quote(GLINT_TYPE)}")
    bedrijf_ids = {b["id"] for b in bedrijven}

    contacten = haal(url, key, "contacts",
                     "id,full_name,first_name,last_name,title,email,email_status,linkedin_url,"
                     "country,stage,former,do_not_email,marketing_content_opt_in,last_email_date,"
                     "company_id,company_name,owner,source")
    outreach = haal(url, key, "outreach_contact", "email,status,company")

    outreach_op_mail = {(o.get("email") or "").lower(): o for o in outreach if o.get("email")}

    per_bedrijf = defaultdict(list)
    for c in contacten:
        if c.get("company_id") in bedrijf_ids:
            per_bedrijf[c["company_id"]].append(c)

    wb = openpyxl.Workbook()

    # --- Bedrijven -----------------------------------------------------------
    ws = wb.active
    ws.title = "Bedrijven"
    kol_b = ["Account", "Bedrijf", "Land", "Stad", "Branche", "Medewerkers", "Website",
             "Contactpersonen", "Met e-mail", "Benaderbaar", "Al benaderd", "Stage", "Eigenaar"]
    rijen_b = []
    for b in sorted(bedrijven, key=lambda x: (x.get("name") or "").lower()):
        cs = per_bedrijf.get(b["id"], [])
        met_mail = [c for c in cs if (c.get("email") or "").strip()]
        kan = [c for c in cs if benaderbaar(c)[0] == "ja"]
        al = [c for c in cs if (c.get("email") or "").lower() in outreach_op_mail]
        rijen_b.append([b.get("account_no"), b.get("name"), b.get("country"), b.get("city"),
                        b.get("industry") or b.get("sub_industry"), b.get("employee_count"),
                        b.get("website"), len(cs), len(met_mail), len(kan), len(al),
                        b.get("stage"), b.get("owner")])
    kop(ws, kol_b)
    for r in rijen_b:
        ws.append(r)
    breedtes(ws, kol_b, rijen_b)

    # --- Contactpersonen -----------------------------------------------------
    ws = wb.create_sheet("Contactpersonen")
    kol_c = ["Bedrijf", "Naam", "Functie", "E-mail", "Adresbron", "Benaderbaar", "Reden",
             "Al benaderd", "Outreach-status", "Opt-in marketing", "Laatste mail",
             "LinkedIn", "Land", "Eigenaar"]
    naam_van_bedrijf = {b["id"]: b.get("name") for b in bedrijven}
    rijen_c = []
    for bid, cs in per_bedrijf.items():
        for c in sorted(cs, key=lambda x: (x.get("last_name") or "").lower()):
            kan, reden = benaderbaar(c)
            o = outreach_op_mail.get((c.get("email") or "").lower())
            rijen_c.append([
                naam_van_bedrijf.get(bid),
                c.get("full_name") or f"{c.get('first_name') or ''} {c.get('last_name') or ''}".strip(),
                c.get("title"), c.get("email"), c.get("email_status"), kan, reden,
                "ja" if o else "", (o or {}).get("status", ""),
                "ja" if c.get("marketing_content_opt_in") else "",
                (c.get("last_email_date") or "")[:10], c.get("linkedin_url"),
                c.get("country"), c.get("owner")])
    rijen_c.sort(key=lambda r: ((r[0] or "").lower(), (r[1] or "").lower()))
    kop(ws, kol_c)
    for r in rijen_c:
        ws.append(r)
    breedtes(ws, kol_c, rijen_c)

    # --- Marktanalyse van de collega ----------------------------------------
    rijen_a = []
    kol_a = []
    if os.path.exists(ANALYSE_XLSX):
        awb = openpyxl.load_workbook(ANALYSE_XLSX, read_only=True, data_only=True)
        aws = awb["Prospects"] if "Prospects" in awb.sheetnames else awb.worksheets[0]
        # Het blad opent met een titel en een blok totalen; de echte kolomkoppen
        # staan pas op de regel die met 'Company' begint. Rij 1 als kop nemen
        # levert een tabel met lege koppen op die er wel uitziet alsof hij klopt.
        alle = list(aws.iter_rows(values_only=True))
        start = next((i for i, r in enumerate(alle)
                      if str((r[0] or "")).strip().lower() == "company"), 0)
        kop_a = [str(x or "").strip() for x in alle[start]]
        bekend = {norm(b.get("name")) for b in bedrijven}
        adres_op_naam = {}
        for c in contacten:
            n = norm(c.get("full_name") or
                     f"{c.get('first_name') or ''} {c.get('last_name') or ''}")
            if n and c.get("email"):
                adres_op_naam.setdefault(n, c["email"])
        kol_a = kop_a + ["Bedrijf in CRM", "Contact in CRM", "E-mailadres in CRM"]
        for rij in alle[start + 1:]:
            if not any(rij):
                continue
            d = dict(zip(kop_a, rij))
            bedrijf = next((d.get(k) for k in ("Company", "Company Name", "Bedrijf", "Account")
                            if d.get(k)), "")
            persoon = next((d.get(k) for k in ("Primary contact candidate",
                                               "Primary contact candidates", "Contact")
                            if d.get(k)), "")
            mail = adres_op_naam.get(norm(persoon))
            rijen_a.append(list(rij) + [
                "ja" if norm(bedrijf) in bekend else "nee",
                "ja" if mail else "nee", mail or ""])
        ws = wb.create_sheet("Marktanalyse collega")
        kop(ws, kol_a)
        for r in rijen_a:
            ws.append(r)
        breedtes(ws, kol_a, rijen_a)

    # --- Surfe-ronde Prioriteit A -------------------------------------------
    rijen_s = []
    if os.path.exists(SURFE_IN) and os.path.exists(SURFE_UIT):
        inp = list(csv.DictReader(open(SURFE_IN)))
        uit = list(csv.DictReader(open(SURFE_UIT)))

        def li(u):
            m = re.search(r"linkedin\.com/in/([^/?]+)", (u or "").lower())
            return m.group(1) if m else None

        op_li = {li(r["LinkedIn URL"]): r for r in inp if li(r["LinkedIn URL"])}
        op_naam = {norm(r["First Name"]) + "|" + norm(r["Last Name"]): r for r in inp}
        in_crm = {(c.get("email") or "").lower() for c in contacten if c.get("email")}
        for o in uit:
            src = op_li.get(li(o["LinkedIn URL"])) or op_naam.get(
                norm(o["First Name"]) + "|" + norm(o["Last Name"])) or {}
            oud = (src.get("Email") or "").strip().lower()
            nieuw = (o.get("Email") or "").strip().lower()
            if not nieuw:
                uitkomst = "geen adres gevonden"
            elif not oud:
                uitkomst = "nieuw adres"
            elif oud == nieuw:
                uitkomst = "bevestigd"
            elif oud.split("@")[0] == nieuw.split("@")[0]:
                uitkomst = "domein gecorrigeerd"
            else:
                uitkomst = "ander adres"
            rijen_s.append([
                f"{src.get('First Name', o['First Name'])} {src.get('Last Name', o['Last Name'])}".strip(),
                src.get("Company Name") or o.get("Company Name"), src.get("Job Title") or o.get("Job Title"),
                src.get("Prioriteit"), src.get("Herkomst"), src.get("Bewijsniveau"),
                oud, nieuw, o.get("Email Validation Status"), uitkomst,
                "ja" if nieuw and nieuw in in_crm else "nee",
                o.get("Current Employer Name"), src.get("LinkedIn URL") or o.get("LinkedIn URL")])
        kol_s = ["Naam", "Bedrijf", "Functie", "Prioriteit", "Herkomst", "Bewijsniveau",
                 "Adres vooraf", "Adres via Surfe", "Validatie", "Uitkomst",
                 "Staat in CRM", "Werkgever volgens Surfe", "LinkedIn"]
        rijen_s.sort(key=lambda r: ((r[9] or ""), (r[1] or "").lower()))
        ws = wb.create_sheet("Surfe Prioriteit A")
        kop(ws, kol_s)
        for r in rijen_s:
            ws.append(r)
        breedtes(ws, kol_s, rijen_s)

    # --- Overzicht vooraan ---------------------------------------------------
    ws = wb.create_sheet("Overzicht", 0)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 14
    ws["A1"] = "Glint-prospects, totaaloverzicht"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Gemaakt op {date.today():%d-%m-%Y}"
    ws["A2"].font = Font(italic=True, color="666666")

    kan_totaal = sum(1 for r in rijen_c if r[5] == "ja")
    met_mail = sum(1 for r in rijen_c if r[3])
    al_benaderd = sum(1 for r in rijen_c if r[7] == "ja")
    regels = [
        ("In het CRM", None),
        ("Accounts met type 'Expected Glint Customers'", len(bedrijven)),
        ("Waarvan met minstens een contactpersoon", sum(1 for b in bedrijven if per_bedrijf.get(b["id"]))),
        ("Contactpersonen bij die accounts", len(rijen_c)),
        ("Waarvan met e-mailadres", met_mail),
        ("Waarvan vandaag benaderbaar", kan_totaal),
        ("Waarvan al in een outreachlijst", al_benaderd),
        ("", None),
        ("Marktanalyse collega", None),
        ("Rijen in het analysebestand", len(rijen_a)),
        ("Waarvan het bedrijf al als Glint-account bestaat", sum(1 for r in rijen_a if r[-3] == "ja")),
        ("Waarvan de genoemde persoon al in het CRM staat", sum(1 for r in rijen_a if r[-2] == "ja")),
        ("", None),
        ("Surfe-validatie Prioriteit A", None),
        ("Aangeboden ter validatie", len(rijen_s)),
        ("Adres bevestigd", sum(1 for r in rijen_s if r[9] == "bevestigd")),
        ("Nieuw adres gevonden", sum(1 for r in rijen_s if r[9] == "nieuw adres")),
        ("Ander adres (oude was fout)", sum(1 for r in rijen_s if r[9] == "ander adres")),
        ("Domein gecorrigeerd", sum(1 for r in rijen_s if r[9] == "domein gecorrigeerd")),
        ("Geen adres gevonden", sum(1 for r in rijen_s if r[9] == "geen adres gevonden")),
    ]
    rij = 4
    for label, waarde in regels:
        if waarde is None and label:
            ws.cell(row=rij, column=1, value=label).font = Font(bold=True, color="1F3864")
        elif label:
            ws.cell(row=rij, column=1, value=label)
            ws.cell(row=rij, column=2, value=waarde)
        rij += 1

    wb.save(args.out)
    print(f"Geschreven: {args.out}")
    print(f"  Bedrijven {len(rijen_b)} | Contactpersonen {len(rijen_c)} | "
          f"Analyse {len(rijen_a)} | Surfe {len(rijen_s)}")


if __name__ == "__main__":
    main()
