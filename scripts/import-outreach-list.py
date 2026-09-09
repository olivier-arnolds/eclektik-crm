#!/usr/bin/env python3
"""
Importeert de outreach-lijst (xlsx) in Supabase: outreach_campaign + outreach_contact.
Ontwerp en besluiten: docs/outreach-handover.md (§4 en addendum §9).

WAAROM PYTHON EN NIET JS
  Dit is een eenmalig, lokaal admin-script. De app-dependencies blijven zo schoon
  (geen xlsx-library in package.json) en openpyxl staat al op de machine. Het script
  draait LOKAAL en praat direct met Supabase, zodat de prospectgegevens niet via een
  omweg reizen. Daarom staat de xlsx ook bewust niet in git.

GEBRUIK
  Zet eenmalig in .env.local in de repo-root (gitignored, wordt nooit gecommit):
    SUPABASE_URL=https://jdzaypckluncdwsoxurs.supabase.co
    SUPABASE_SERVICE_KEY=<service_role key uit Supabase, Settings > API>

  # 1) Dry-run (standaard, schrijft niets):
  python3 scripts/import-outreach-list.py --file "/pad/naar/masterlijst-final-aangevuld.xlsx"

  # 2) Echt importeren:
  python3 scripts/import-outreach-list.py --file "..." --apply

  # 3) Teksten van bestaande rijen verversen (alleen bewust gebruiken):
  python3 scripts/import-outreach-list.py --file "..." --apply --overwrite

De service_role key haal je uit Supabase (Settings, API) of uit de Vercel-env.
Zonder --apply gebeurt er niets: je krijgt alleen het rapport.

WAT HET DOET
  * Leest de xlsx en splitst elke berichtkolom in subject + body ("Subject: ..." op
    regel 1, lege regel, dan de body).
  * Leidt priority_tier af uit het kleurprefix (groen/geel/oranje -> top/good/medium).
  * Rijen met 'Reserve' in Outreach-prio komen mee als is_reserve=true en status
    'paused', zodat ze nooit per ongeluk in golf 1 meelopen.
  * Cross-match op e-mail en domein tegen contacts en companies, zodat bestaande
    klantrelaties en do_not_email automatisch op 'paused' komen.
  * Upsert per 100 rijen. Een tweede run voegt standaard alleen NIEUWE rijen toe en
    raakt bestaande niet aan (anders zou je live campagnevoortgang wissen).
"""

import argparse
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl ontbreekt. Installeer met: python3 -m pip install openpyxl")

# Bedrijfstypes die duiden op een lopende relatie. Een koude uitnodiging gaat daar
# niet automatisch heen; die zet het script op 'paused' zodat jij eerst afstemt.
# Pas deze set aan als de afspraak verandert.
PAUSE_COMPANY_TYPES = {"Customer", "Partner", "Relation"}

BATCH = 100
TIER_BY_PREFIX = {"\U0001F7E2": "top", "\U0001F7E1": "good", "\U0001F7E0": "medium"}


# ── helpers ────────────────────────────────────────────────────────────────────
def load_env_local():
    """Vult ontbrekende env-vars uit .env.local in de repo-root.

    Zo hoef je geen service-key in je shell-history te zetten. .env.local is
    gitignored (.env*.local) en wordt nooit gecommit. Bestaande env-vars winnen.
    """
    path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env.local")
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v


def s(v):
    """Cel naar getrimde string, of '' als leeg."""
    return "" if v is None else str(v).strip()


def email_domain(email):
    return email.split("@", 1)[1].lower() if "@" in email else None


def norm_name(v):
    return " ".join(s(v).lower().split())


def domain_from_website(website):
    w = s(website).lower()
    if not w:
        return None
    if "//" not in w:
        w = "//" + w
    host = urllib.parse.urlparse(w).netloc or ""
    host = host.split("@")[-1].split(":")[0]
    return host[4:] if host.startswith("www.") else (host or None)


def tier_from_label(label):
    lab = s(label)
    for prefix, tier in TIER_BY_PREFIX.items():
        if lab.startswith(prefix):
            return tier
    return None


def split_subject_body(cell):
    """'Subject: X\\n\\nbody...' -> ('X', 'body...'). Zonder Subject-regel: (None, tekst)."""
    text = s(cell)
    if not text:
        return None, None
    lines = text.splitlines()
    if lines and lines[0].lower().startswith("subject:"):
        subject = lines[0].split(":", 1)[1].strip()
        rest = lines[1:]
        while rest and not rest[0].strip():
            rest.pop(0)
        return (subject or None), ("\n".join(rest).strip() or None)
    return None, text


def find_col(headers, *prefixes):
    """Kolomindex op (case-insensitive) prefix; de berichtkolommen hebben lange namen."""
    low = [h.lower() for h in headers]
    for p in prefixes:
        pl = p.lower()
        for i, h in enumerate(low):
            if h == pl:
                return i
        for i, h in enumerate(low):
            if h.startswith(pl):
                return i
    return None


# ── Supabase REST (stdlib) ────────────────────────────────────────────────────
class Supa:
    def __init__(self, url, key):
        self.base = url.rstrip("/") + "/rest/v1"
        self.key = key

    def _req(self, method, path, body=None, extra_headers=None):
        headers = {
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        }
        headers.update(extra_headers or {})
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                raw = r.read().decode() or "[]"
                return json.loads(raw) if raw.strip().startswith(("[", "{")) else []
        except urllib.error.HTTPError as e:
            sys.exit(f"Supabase {method} {path} faalde ({e.code}): {e.read().decode()[:500]}")

    def select_all(self, table, columns):
        """Alles ophalen met paginering (PostgREST geeft standaard max 1000)."""
        out, offset = [], 0
        while True:
            q = f"/{table}?select={urllib.parse.quote(columns)}&limit=1000&offset={offset}"
            page = self._req("GET", q)
            out.extend(page)
            if len(page) < 1000:
                return out
            offset += 1000

    def insert(self, table, rows, on_conflict=None, overwrite=False):
        path = f"/{table}"
        headers = {"Prefer": "return=representation"}
        if on_conflict:
            path += f"?on_conflict={urllib.parse.quote(on_conflict)}"
            headers["Prefer"] = (
                "resolution=merge-duplicates,return=representation" if overwrite
                else "resolution=ignore-duplicates,return=representation"
            )
        return self._req("POST", path, rows, headers)


# ── main ──────────────────────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser(description="Importeer de outreach-lijst in Supabase.")
    ap.add_argument("--file", required=True, help="pad naar de xlsx")
    ap.add_argument("--campaign-name", default="Amsterdam 2026")
    ap.add_argument("--sender", default="marco@eclectik.co")
    ap.add_argument("--hard-stop", default="2026-10-02T00:00:00Z")
    ap.add_argument("--daily-cap", type=int, default=60, help="startcap; ramp naar 120")
    ap.add_argument("--apply", action="store_true", help="schrijf naar de database")
    ap.add_argument("--overwrite", action="store_true",
                    help="werk ook BESTAANDE rijen bij (wist campagnevoortgang)")
    args = ap.parse_args()

    load_env_local()
    url = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit(
            "SUPABASE_URL en SUPABASE_SERVICE_KEY ontbreken.\n"
            "Zet ze in .env.local in de repo-root (gitignored), bijvoorbeeld:\n"
            "  SUPABASE_URL=https://jdzaypckluncdwsoxurs.supabase.co\n"
            "  SUPABASE_SERVICE_KEY=<service_role key uit Supabase, Settings > API>\n"
            "of geef ze mee als env-var op de commandoregel."
        )
    supa = Supa(url, key)

    # 1. Lijst lezen
    wb = openpyxl.load_workbook(args.file, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    headers = [s(h) for h in next(it)]
    rows = [r for r in it if any(c not in (None, "") for c in r)]
    wb.close()

    C = {
        "first": find_col(headers, "Voornaam"),
        "last": find_col(headers, "Achternaam"),
        "title": find_col(headers, "Titel"),
        "linkedin": find_col(headers, "LinkedIn"),
        "email": find_col(headers, "Find work email", "Work Email", "Email"),
        "prio_label": find_col(headers, "Prioriteit"),
        "prio_num": find_col(headers, "Outreach-prio"),
        "source": find_col(headers, "Bron"),
        "hook": find_col(headers, "Uitnodigingshaak"),
        "company": find_col(headers, "Bedrijf"),
        "website": find_col(headers, "Website"),
        "location": find_col(headers, "Locatie"),
        "employees": find_col(headers, "Employee Count"),
        "size": find_col(headers, "Size"),
        "industry": find_col(headers, "Industry"),
        "company_li": find_col(headers, "Company Linkedin", "Compamy_Linkedin"),
        "msg1": find_col(headers, "Message 1"),
        "msg2": find_col(headers, "Message 2"),
        "frontline": find_col(headers, "Frontline workers?"),
    }
    missing = [k for k in ("email", "prio_label", "prio_num", "msg1", "msg2") if C[k] is None]
    if missing:
        sys.exit(f"Kolommen niet gevonden: {missing}\nHeaders in het bestand: {headers}")

    g = lambda r, k: s(r[C[k]]) if (C[k] is not None and C[k] < len(r)) else ""

    # 2. CRM inlezen voor de cross-match
    print("CRM inlezen voor de cross-match ...")
    contacts = supa.select_all("contacts", "id,email,company_id,do_not_email")
    companies = supa.select_all("companies", "id,name,website,type")
    c_by_email = {c["email"].strip().lower(): c for c in contacts if c.get("email")}
    co_by_id = {c["id"]: c for c in companies}
    co_by_domain, co_by_name = {}, {}
    for co in companies:
        d = domain_from_website(co.get("website"))
        if d:
            co_by_domain.setdefault(d, co)
        n = norm_name(co.get("name"))
        if n:
            co_by_name.setdefault(n, co)
    print(f"  {len(contacts)} contacts, {len(companies)} companies ingelezen")

    # 3. Rijen omzetten
    now_iso = datetime.now(timezone.utc).isoformat()
    out, skipped_no_email, seen = [], 0, set()
    stats = {"reserve": 0, "queued": 0, "paused": 0, "match_contact": 0, "match_company": 0}
    pause_reasons, relation_hits = {}, []

    for r in rows:
        email = g(r, "email").lower()
        if not email or "@" not in email:
            skipped_no_email += 1
            continue
        if email in seen:          # dubbel adres in het bestand: eerste wint
            continue
        seen.add(email)

        dom = email_domain(email)
        prio_raw = g(r, "prio_num")
        is_num = prio_raw.replace(".0", "").isdigit()
        is_reserve = not is_num
        m1_sub, m1_body = split_subject_body(r[C["msg1"]] if C["msg1"] < len(r) else None)
        m2_sub, m2_body = split_subject_body(r[C["msg2"]] if C["msg2"] < len(r) else None)

        contact = c_by_email.get(email)
        company = None
        if contact and contact.get("company_id"):
            company = co_by_id.get(contact["company_id"])
        if company is None and dom:
            company = co_by_domain.get(dom)
        if company is None:
            company = co_by_name.get(norm_name(g(r, "company")))
        if contact:
            stats["match_contact"] += 1
        if company:
            stats["match_company"] += 1

        # Statusbepaling. Zwaarste reden eerst, want de reden wordt vastgelegd.
        status, reason = "queued", None
        if contact and contact.get("do_not_email"):
            status, reason = "paused", "do_not_email in CRM"
        elif company and s(company.get("type")) in PAUSE_COMPANY_TYPES:
            status = "paused"
            reason = f"bestaande relatie ({company['type']}): eerst afstemmen"
            relation_hits.append((g(r, "company") or company.get("name"), company["type"], email))
        elif is_reserve:
            status, reason = "paused", "reserve (geen nummer in Outreach-prio)"
        elif not (m1_body and m2_body):
            status, reason = "paused", "teksten ontbreken"

        if is_reserve:
            stats["reserve"] += 1
        stats[status] += 1
        if reason:
            pause_reasons[reason] = pause_reasons.get(reason, 0) + 1

        out.append({
            "first_name": g(r, "first") or None,
            "last_name": g(r, "last") or None,
            "title": g(r, "title") or None,
            "email": email,
            "email_domain": dom,
            "company": g(r, "company") or None,
            "company_linkedin": g(r, "company_li") or None,
            "website": g(r, "website") or None,
            "industry": g(r, "industry") or None,
            "location": g(r, "location") or None,
            "employee_count": g(r, "employees") or None,
            "size_bucket": g(r, "size") or None,
            "linkedin_url": g(r, "linkedin") or None,
            "priority_label": g(r, "prio_label") or None,
            "priority_tier": tier_from_label(g(r, "prio_label")),
            "outreach_prio": int(prio_raw.replace(".0", "")) if is_num else None,
            "is_reserve": is_reserve,
            "hook_note": g(r, "hook") or None,
            "frontline_note": g(r, "frontline") or None,
            "source": g(r, "source") or None,
            "msg1_subject": m1_sub, "msg1_body": m1_body,
            "msg2_subject": m2_sub, "msg2_body": m2_body,
            "status": status,
            "paused_reason": reason,
            "next_action_at": now_iso if status == "queued" else None,
            "contact_id": contact["id"] if contact else None,
            "company_id": company["id"] if company else None,
        })

    # 4. Rapport
    wave = [o for o in out if not o["is_reserve"]]
    ready = [o for o in wave if o["status"] == "queued"]
    print(f"\n{'=' * 62}\nRAPPORT ({'DRY-RUN, er wordt niets geschreven' if not args.apply else 'APPLY'})\n{'=' * 62}")
    print(f"  rijen in bestand      : {len(rows)}")
    print(f"  zonder e-mail overgeslagen: {skipped_no_email}")
    print(f"  te importeren         : {len(out)}")
    print(f"  waarvan reserve       : {stats['reserve']}")
    print(f"  golf 1 (met nummer)   : {len(wave)}")
    print(f"  DIRECT VERZENDBAAR    : {len(ready)}")
    print(f"  status queued/paused  : {stats['queued']} / {stats['paused']}")
    print(f"  gematcht op CRM-contact: {stats['match_contact']}   op bedrijf: {stats['match_company']}")
    tiers = {}
    for o in wave:
        tiers[o["priority_tier"] or "?"] = tiers.get(o["priority_tier"] or "?", 0) + 1
    print(f"  tiers in golf 1       : {tiers}")
    if pause_reasons:
        print("\n  redenen voor paused:")
        for k, v in sorted(pause_reasons.items(), key=lambda x: -x[1]):
            print(f"    {v:5d}  {k}")
    if relation_hits:
        print(f"\n  !! {len(relation_hits)} prospects bij een BESTAANDE relatie (op paused gezet):")
        for comp, typ, em in relation_hits[:40]:
            print(f"    - {comp} ({typ})  {em}")
        if len(relation_hits) > 40:
            print(f"    ... en {len(relation_hits) - 40} meer")

    if not args.apply:
        print("\nDry-run klaar. Voeg --apply toe om echt te importeren.")
        return

    # 5. Campagne find-or-create
    existing = supa._req("GET", f"/outreach_campaign?select=id,name&name=eq."
                                f"{urllib.parse.quote(args.campaign_name)}&limit=1")
    if existing:
        campaign_id = existing[0]["id"]
        print(f"\nBestaande campagne hergebruikt: {args.campaign_name} ({campaign_id})")
    else:
        created = supa.insert("outreach_campaign", [{
            "name": args.campaign_name,
            "sender_mailbox": args.sender,
            "daily_cap": args.daily_cap,
            "hard_stop_at": args.hard_stop,
            "status": "draft",
        }])
        campaign_id = created[0]["id"]
        print(f"\nCampagne aangemaakt: {args.campaign_name} ({campaign_id})")

    # 6. Upsert in batches
    for o in out:
        o["campaign_id"] = campaign_id
    written = 0
    for i in range(0, len(out), BATCH):
        chunk = out[i:i + BATCH]
        res = supa.insert("outreach_contact", chunk,
                          on_conflict="campaign_id,email", overwrite=args.overwrite)
        written += len(res)
        print(f"  batch {i // BATCH + 1}: {len(res)} van {len(chunk)} weggeschreven")
    print(f"\nKlaar. {written} rijen weggeschreven"
          f"{' (bestaande bijgewerkt)' if args.overwrite else ' (bestaande overgeslagen)'}.")
    print("Controleer met: select status, count(*) from outreach_contact group by status;")


if __name__ == "__main__":
    main()
