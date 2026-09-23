#!/usr/bin/env python3
"""
Lijst van Prioriteit A-bedrijven uit de marktanalyse waar we (bijna) niemand van
kennen, zodat er gericht namen bij gezocht kunnen worden.

Drie groepen, want ze vragen een ander soort werk:
  1. geen contact in het CRM               -> naam zoeken
  2. wel contacten, maar geen enkel adres  -> adres zoeken
  3. wel contacten met adres, maar niemand senior -> betere naam zoeken

Die derde groep is er omdat de rolbeoordeling van de collega liet zien dat een
bedrijf met vijf recruiters in het CRM er in de praktijk net zo voor staat als
een bedrijf met nul contacten.

GEBRUIK
  python3 scripts/prio-a-zonder-contacten.py [--excel PAD] [--out PAD]
"""
import argparse, csv, json, os, re, sys, unicodedata
import urllib.parse, urllib.request
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Titels die wijzen op iemand die over employee listening beslist. Bewust ruim:
# de lijst is om te sorteren, niet om iemand af te wijzen.
SENIOR = re.compile(
    r'\b(chief|chro|cpo|c-level|evp|svp|vp|vice[- ]president|head of|global head|'
    r'director|people analytics|employee (listening|experience|engagement))\b', re.I)


def env():
    w = {}
    with open(os.path.join(ROOT, '.env.local')) as f:
        for r in f:
            if '=' in r and not r.strip().startswith('#'):
                k, _, v = r.partition('=')
                w[k.strip()] = v.strip().strip('"')
    if not w.get('SUPABASE_URL') or not w.get('SUPABASE_SERVICE_KEY'):
        sys.exit('SUPABASE_URL of SUPABASE_SERVICE_KEY ontbreekt in .env.local')
    return w['SUPABASE_URL'].rstrip('/'), w['SUPABASE_SERVICE_KEY']


def haal(url, key, tabel, kolommen):
    uit, offset = [], 0
    while True:
        pad = f'{url}/rest/v1/{tabel}?select={urllib.parse.quote(kolommen)}&limit=1000&offset={offset}'
        req = urllib.request.Request(pad, headers={
            'apikey': key, 'Authorization': f'Bearer {key}', 'Accept': 'application/json'})
        with urllib.request.urlopen(req, timeout=120) as r:
            blok = json.loads(r.read().decode())
        uit.extend(blok)
        if len(blok) < 1000:
            return uit
        offset += 1000


def sleutel(s):
    """Bedrijfsnamen uit de analyse en uit het CRM lopen uiteen ('Merck & Co.
    (MSD)' tegenover 'Merck & Co'), dus vergelijken op een gestripte sleutel en
    niet op de letterlijke naam."""
    s = unicodedata.normalize('NFKD', (s or '')).encode('ascii', 'ignore').decode()
    s = re.sub(r'\(.*?\)', '', s)
    s = re.sub(r'\b(inc|ltd|plc|bv|nv|ag|sa|gmbh|corporation|corp|group|holding|holdings|limited|co)\b',
               '', s, flags=re.I)
    return re.sub(r'[^a-z0-9]', '', s.lower())


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--excel', default=os.path.expanduser(
        '~/Downloads/glint-prospects-totaaloverzicht-2026-09-22 (1).xlsx'))
    p.add_argument('--out', default=os.path.expanduser(
        f'~/Downloads/prio-a-zonder-contacten-{date.today():%Y-%m-%d}.csv'))
    args = p.parse_args()
    url, key = env()

    bedrijven = haal(url, key, 'companies', 'id,name,website,type,country')
    contacten = haal(url, key, 'contacts',
                     'id,full_name,title,email,company_id,stage,former,do_not_email')

    per_bedrijf = {}
    for b in bedrijven:
        per_bedrijf.setdefault(sleutel(b['name']), []).append(b)
    contacten_per_id = {}
    for c in contacten:
        if (c.get('stage') or '').lower() == 'inactive' or c.get('former'):
            continue
        contacten_per_id.setdefault(c.get('company_id'), []).append(c)

    wb = openpyxl.load_workbook(args.excel, data_only=True)
    ws = wb['Marktanalyse collega']
    kop = [c.value for c in ws[1]]
    K = {k: i for i, k in enumerate(kop)}

    rijen = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r) or str(r[K['Priority']] or '').strip().upper() != 'A':
            continue
        naam = r[K['Company']]
        crm = per_bedrijf.get(sleutel(naam), [])
        cs = [c for b in crm for c in contacten_per_id.get(b['id'], [])]
        met_adres = [c for c in cs if (c.get('email') or '').strip()]
        senior = [c for c in met_adres if SENIOR.search(c.get('title') or '')]

        if not cs:
            groep, wat = '1 geen contact', 'naam zoeken'
        elif not met_adres:
            groep, wat = '2 geen adres', 'adres zoeken'
        elif not senior:
            groep, wat = '3 niemand senior', 'betere naam zoeken'
        else:
            continue

        rijen.append({
            'Groep': groep, 'Actie': wat, 'Bedrijf': naam,
            'In CRM als': crm[0]['name'] if crm else '(bestaat niet)',
            'Website': (crm[0].get('website') or '') if crm else '',
            'Land': r[K['Country/market']] or '',
            'Sector': r[K['Sector']] or '',
            'Bewijsniveau': r[K['Evidence level']] or '',
            'Doelrollen volgens analyse': r[K['Target roles']] or '',
            'Kandidaat uit de analyse': r[K['Primary contact candidate']] or '',
            'Contacten in CRM': len(cs),
            'Waarvan met adres': len(met_adres),
            'Huidige namen': ', '.join(f"{c['full_name']} ({c.get('title') or 'geen functie'})" for c in cs[:6]),
        })

    rijen.sort(key=lambda x: (x['Groep'], x['Bedrijf'].lower()))
    kol = list(rijen[0].keys()) if rijen else []
    with open(args.out, 'w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=kol)
        w.writeheader()
        w.writerows(rijen)

    from collections import Counter
    telling = Counter(r['Groep'] for r in rijen)
    print(f'Geschreven: {args.out}')
    print(f'  {len(rijen)} van de 74 Prioriteit A-bedrijven vragen aandacht:')
    for g in sorted(telling):
        print(f'    {g}: {telling[g]}')


if __name__ == '__main__':
    main()
