#!/usr/bin/env python3
"""
Importeert de fase 1 en fase 2 lijst in een bestaande outreach-campagne.

WAAROM EEN EIGEN SCRIPT NAAST import-outreach-list.py
  Dat script verwacht het onderwerp IN de berichtkolom ('Subject: ...' op regel 1,
  lege regel, dan de body). Dit bestand heeft aparte kolommen Onderwerp en Bericht.
  Twee formaten in een script gaan uit elkaar lopen, en het bestaande script draait
  al goed voor de Amsterdam-lijst. Vandaar een eigen importer met dezelfde
  conventies: droge run als standaard, sleutel uit .env.local, rapport voor je
  iets wegschrijft.

WAT HET DOET
  * Staat het adres al in de campagne, dan worden ALLEEN onderwerp, bericht en
    notitie bijgewerkt. Status, tokens en voortgang blijven ongemoeid; die wil je
    niet verliezen bij een tweede run.
  * Staat het adres er nog niet in, dan komt er een nieuwe rij bij, status queued.
  * De kolom 'Let op' gaat samen met de bewijskracht naar hook_note, zodat je in
    de tab ziet waarom iemand in de lijst staat.

GEBRUIK
  Zet eenmalig in .env.local in de repo-root (gitignored):
    SUPABASE_URL=https://jdzaypckluncdwsoxurs.supabase.co
    SUPABASE_SERVICE_KEY=<service_role key>

  # Droge run, schrijft niets:
  python3 scripts/import-outreach-fase1-fase2.py --file "/pad/naar/bestand.xlsx"

  # Echt importeren:
  python3 scripts/import-outreach-fase1-fase2.py --file "..." --apply
"""

import argparse
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request

import openpyxl

CAMPAGNE_STANDAARD = "Glint Prioriteit A"
BATCH = 50


def load_env_local():
    """Vult ontbrekende env-vars uit .env.local in de repo-root.

    Zo hoeft de service-key niet in je shell-history. Bestaande env-vars winnen.
    """
    path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env.local")
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, v = line.split("=", 1)
            k, v = k.strip(), v.strip().strip('"').strip("'")
            if k and k not in os.environ:
                os.environ[k] = v


def tekst(v):
    """Een cel als nette string, of None. Excel geeft None, getallen en spaties."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


class Supabase:
    def __init__(self, url, key):
        self.base = url.rstrip("/") + "/rest/v1"
        self.key = key

    def _req(self, method, path, body=None, extra_headers=None):
        headers = {
            "apikey": self.key,
            "Authorization": f"Bearer {self.key}",
            "Content-Type": "application/json",
            "Accept": "application/json",
        }
        headers.update(extra_headers or {})
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(self.base + path, data=data, headers=headers, method=method)
        try:
            with urllib.request.urlopen(req, timeout=180) as r:
                raw = r.read().decode() or "[]"
                return json.loads(raw) if raw.strip().startswith(("[", "{")) else []
        except urllib.error.HTTPError as e:
            sys.exit(f"Supabase {method} {path} faalde ({e.code}): {e.read().decode()[:500]}")

    def select_all(self, table, columns, filter_q=""):
        out, offset = [], 0
        while True:
            q = f"/{table}?select={urllib.parse.quote(columns)}{filter_q}&limit=1000&offset={offset}"
            page = self._req("GET", q)
            out.extend(page)
            if len(page) < 1000:
                return out
            offset += 1000

    def insert(self, table, rows):
        return self._req("POST", f"/{table}", rows, {"Prefer": "return=representation"})

    def patch(self, table, filter_q, body):
        return self._req("PATCH", f"/{table}?{filter_q}", body, {"Prefer": "return=representation"})


def notitie_van(rij):
    """De kolom 'Let op' plus de bewijskracht, als een leesbare regel.

    Bewust samengevoegd in hook_note en niet over losse kolommen verdeeld: het is
    context voor wie de lijst naloopt, geen data waar het systeem op stuurt.
    """
    delen = []
    for sleutel in ("Let op", "Bewijskracht", "Soort bewijs"):
        v = tekst(rij.get(sleutel))
        if not v:
            continue
        delen.append(v if sleutel == "Let op" else f"{sleutel.lower()}: {v}")
    return " | ".join(delen) or None


def lees_xlsx(pad):
    wb = openpyxl.load_workbook(pad, data_only=True)
    ws = wb.worksheets[0]
    kop = [c.value for c in ws[1]]
    verplicht = {"Naam", "E-mail", "Onderwerp", "Bericht"}
    ontbreekt = verplicht - set(k for k in kop if k)
    if ontbreekt:
        sys.exit(f"Kolommen ontbreken in het bestand: {', '.join(sorted(ontbreekt))}")
    rijen = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        rijen.append(dict(zip(kop, r)))
    return rijen


def main():
    ap = argparse.ArgumentParser(description="Importeer de fase 1 en 2 lijst in een outreach-campagne.")
    ap.add_argument("--file", required=True, help="pad naar de xlsx")
    ap.add_argument("--campaign-name", default=CAMPAGNE_STANDAARD)
    ap.add_argument("--apply", action="store_true", help="schrijf echt weg (zonder deze vlag gebeurt er niets)")
    args = ap.parse_args()

    load_env_local()
    url = os.environ.get("SUPABASE_URL") or os.environ.get("VITE_SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit("SUPABASE_URL en SUPABASE_SERVICE_KEY ontbreken (zet ze in .env.local).")
    sb = Supabase(url, key)

    camps = sb.select_all("outreach_campaign", "id,name,status,channel",
                          f"&name=eq.{urllib.parse.quote(args.campaign_name)}")
    if not camps:
        sys.exit(f"Campagne '{args.campaign_name}' niet gevonden.")
    camp = camps[0]

    rijen = lees_xlsx(args.file)
    bestaand = sb.select_all("outreach_contact", "id,email",
                             f"&campaign_id=eq.{camp['id']}")
    per_mail = {str(r["email"]).strip().lower(): r["id"] for r in bestaand if r.get("email")}

    bij, nieuw, zonder_mail = [], [], []
    gezien = set()
    for rij in rijen:
        mail = (tekst(rij.get("E-mail")) or "").lower()
        if not mail or "@" not in mail:
            zonder_mail.append(tekst(rij.get("Naam")) or "(naamloos)")
            continue
        if mail in gezien:
            continue
        gezien.add(mail)

        velden = {
            "msg1_subject": tekst(rij.get("Onderwerp")),
            "msg1_body": tekst(rij.get("Bericht")),
            "hook_note": notitie_van(rij),
        }
        if mail in per_mail:
            bij.append((per_mail[mail], mail, velden))
        else:
            naam = tekst(rij.get("Naam")) or ""
            voor, _, achter = naam.partition(" ")
            nieuw.append({
                "campaign_id": camp["id"],
                "first_name": voor or None,
                "last_name": achter or None,
                "email": mail,
                "email_domain": mail.split("@")[-1],
                "title": tekst(rij.get("Functie")),
                "company": tekst(rij.get("Bedrijf")),
                "source": f"Glint {tekst(rij.get('Bron')) or 'import'}",
                "status": "queued",
                "channel": "email",
                **velden,
            })

    print(f"Campagne : {camp['name']} ({camp['status']}, {camp['channel']})")
    print(f"Bestand  : {len(rijen)} rijen")
    print(f"Bijwerken: {len(bij)} (staan al in de campagne, alleen tekst en notitie)")
    print(f"Nieuw    : {len(nieuw)}")
    if zonder_mail:
        print(f"Overgeslagen zonder bruikbaar adres: {len(zonder_mail)} ({', '.join(zonder_mail[:5])})")
    print(f"Campagne na afloop: {len(bestaand)} + {len(nieuw)} = {len(bestaand) + len(nieuw)}")

    if not args.apply:
        print("\nDroge run. Niets weggeschreven. Draai opnieuw met --apply.")
        return

    for i, (rij_id, mail, velden) in enumerate(bij, 1):
        sb.patch("outreach_contact", f"id=eq.{rij_id}", velden)
        if i % 20 == 0:
            print(f"  bijgewerkt: {i}/{len(bij)}")
    print(f"  bijgewerkt: {len(bij)}/{len(bij)}")

    for i in range(0, len(nieuw), BATCH):
        sb.insert("outreach_contact", nieuw[i:i + BATCH])
        print(f"  aangemaakt: {min(i + BATCH, len(nieuw))}/{len(nieuw)}")

    print("\nKlaar.")


if __name__ == "__main__":
    main()
